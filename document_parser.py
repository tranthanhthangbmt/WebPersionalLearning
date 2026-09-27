# document_parser.py — Multi-format Document Text Extractor
"""
Trích xuất text từ nhiều định dạng tài liệu (.txt, .pdf, .docx)
để gửi cho Gemini AI phân tích thành Cây Tri Thức.

Sử dụng:
    from document_parser import parse_document
    text = parse_document("bai_giang.pdf", file_bytes)
"""

import io
import os


# ═══════════════════════════════════════════════════════════
#  PARSERS CHO TỪNG ĐỊNH DẠNG
# ═══════════════════════════════════════════════════════════

def parse_txt(content_bytes: bytes) -> str:
    """Đọc file .txt, thử UTF-8 trước, fallback sang latin-1"""
    for encoding in ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252']:
        try:
            return content_bytes.decode(encoding)
        except (UnicodeDecodeError, AttributeError):
            continue
    # Last resort: decode with replacement
    return content_bytes.decode('utf-8', errors='replace')


def parse_pdf(content_bytes: bytes) -> str:
    """Đọc file .pdf bằng pypdf, trích xuất text từ tất cả các trang"""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ImportError(
            "Thư viện 'pypdf' chưa được cài đặt. "
            "Chạy: pip install pypdf"
        )

    reader = PdfReader(io.BytesIO(content_bytes))
    pages_text = []

    for page_num, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text()
            if text and text.strip():
                pages_text.append(text.strip())
        except Exception as e:
            print(f"⚠️ Lỗi đọc trang {page_num}: {e}")
            continue

    if not pages_text:
        raise ValueError(
            "Không trích xuất được text từ file PDF. "
            "File có thể là ảnh scan — hãy thử file PDF có text layer."
        )

    return "\n\n".join(pages_text)


def parse_docx(content_bytes: bytes) -> str:
    """Đọc file .docx bằng python-docx, trích xuất text từ paragraphs + tables"""
    try:
        from docx import Document
    except ImportError:
        raise ImportError(
            "Thư viện 'python-docx' chưa được cài đặt. "
            "Chạy: pip install python-docx"
        )

    doc = Document(io.BytesIO(content_bytes))
    parts = []

    # Đọc paragraphs (nội dung chính)
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            # Giữ heading style nếu có
            if para.style and para.style.name and 'Heading' in para.style.name:
                parts.append(f"\n{text}\n")
            else:
                parts.append(text)

    # Đọc tables (thường chứa dữ liệu quan trọng)
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_rows.append(" | ".join(cells))
        if table_rows:
            parts.append("\n".join(table_rows))

    if not parts:
        raise ValueError(
            "Không trích xuất được text từ file DOCX. "
            "File có thể rỗng hoặc chỉ chứa hình ảnh."
        )

    return "\n".join(parts)


# ═══════════════════════════════════════════════════════════
#  HÀM CHÍNH: AUTO-DETECT VÀ PARSE
# ═══════════════════════════════════════════════════════════

# Mapping extension -> parser function
_PARSERS = {
    '.txt': parse_txt,
    '.pdf': parse_pdf,
    '.docx': parse_docx,
}

# Các extension được hỗ trợ (dùng cho UI)
SUPPORTED_EXTENSIONS = list(_PARSERS.keys())


def parse_document(filename: str, content_bytes: bytes, max_chars: int = 15000) -> str:
    """
    Trích xuất text từ file tài liệu.

    Args:
        filename: Tên file gốc (dùng để detect extension)
        content_bytes: Nội dung file dạng bytes
        max_chars: Giới hạn ký tự tối đa (default 15,000)

    Returns:
        str: Nội dung text đã trích xuất, cắt tại max_chars

    Raises:
        ValueError: Nếu extension không được hỗ trợ hoặc file rỗng
        ImportError: Nếu thư viện cần thiết chưa cài
    """
    if not filename:
        raise ValueError("Tên file không được để trống.")

    if not content_bytes:
        raise ValueError("Nội dung file rỗng.")

    # Detect extension
    filename = str(filename).strip()
    base_ext = os.path.splitext(filename.lower())[1].strip()
    
    # Normalize: ensure it starts with a dot if it's not empty
    ext = base_ext if (not base_ext or base_ext.startswith('.')) else f'.{base_ext}'
    
    # MAGIC BYTE FALLBACK: If ext is empty or wrong, check file signature
    if not ext or ext not in _PARSERS:
        if content_bytes.startswith(b'%PDF-'):
            ext = '.pdf'
            print(f"🪄 Magic bytes detected: PDF signature found.")
        elif content_bytes.startswith(b'PK\x03\x04') and (b'word/' in content_bytes[:2000] or b'[Content_Types].xml' in content_bytes[:500]):
            ext = '.docx'
            print(f"🪄 Magic bytes detected: DOCX signature found.")
    
    print(f"📄 Parsing document: '{filename}' (final ext: '{ext}')")

    # Fallback for filenames that might not have a dot but are just the extension
    if not ext and filename.lower() in [k.replace('.', '') for k in _PARSERS.keys()]:
        ext = f".{filename.lower()}"

    parser = _PARSERS.get(ext)
    if not parser:
        supported = ", ".join(SUPPORTED_EXTENSIONS)
        raise ValueError(
            f"Định dạng '{ext}' của file '{filename}' không được hỗ trợ. "
            f"Các định dạng hỗ trợ: {supported}"
        )

    # Parse
    text = parser(content_bytes)

    # Clean up
    text = text.strip()
    if not text:
        raise ValueError("File không chứa nội dung text nào.")

    # Truncate
    if len(text) > max_chars:
        # Cắt tại ranh giới câu gần nhất (tránh cắt giữa từ)
        truncated = text[:max_chars]
        last_period = truncated.rfind('.')
        last_newline = truncated.rfind('\n')
        cut_point = max(last_period, last_newline)
        if cut_point > max_chars * 0.8:  # Chỉ cắt nếu điểm cắt hợp lý
            text = truncated[:cut_point + 1]
        else:
            text = truncated
        print(f"📏 Nội dung đã được cắt từ {len(parser(content_bytes))} → {len(text)} ký tự (giới hạn {max_chars})")

    return text


def get_supported_formats_label() -> str:
    """Trả về label hiển thị cho UI (vd: '.txt, .pdf, .docx')"""
    return ", ".join(SUPPORTED_EXTENSIONS)


def get_accept_string() -> str:
    """Trả về chuỗi accept cho HTML upload component (vd: '.txt,.pdf,.docx')"""
    return ",".join(SUPPORTED_EXTENSIONS)
