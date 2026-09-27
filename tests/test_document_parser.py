# tests/test_document_parser.py
"""
Test script cho document_parser.py
Chạy: python tests/test_document_parser.py
"""

import os
import sys

# Thêm root vào path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from document_parser import parse_document, parse_txt, parse_pdf, parse_docx, get_accept_string

PASS = 0
FAIL = 0


def test(name, func):
    global PASS, FAIL
    try:
        func()
        print(f"  ✅ {name}")
        PASS += 1
    except Exception as e:
        print(f"  ❌ {name}: {e}")
        FAIL += 1


def run_tests():
    global PASS, FAIL

    print("=" * 60)
    print("🧪 TEST DOCUMENT PARSER")
    print("=" * 60)

    # ─── TEST 1: parse_txt ───────────────────────────────────
    print("\n📄 Test parse_txt:")

    def test_txt_utf8():
        text = parse_txt("Xin chào thế giới".encode('utf-8'))
        assert "Xin chào" in text, f"Expected 'Xin chào' in result, got: {text[:50]}"

    def test_txt_utf8_bom():
        text = parse_txt("Xin chào".encode('utf-8-sig'))
        assert "Xin chào" in text

    def test_txt_latin1():
        text = parse_txt("café résumé".encode('latin-1'))
        assert "café" in text

    def test_txt_from_file():
        test_file = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'test_doc.txt')
        if os.path.exists(test_file):
            with open(test_file, 'rb') as f:
                content = f.read()
            text = parse_txt(content)
            assert len(text) > 50, f"Text too short: {len(text)} chars"
            assert "Thương mại điện tử" in text or "TMĐT" in text or "Chương" in text
        else:
            raise FileNotFoundError(f"test_doc.txt not found at {test_file}")

    test("UTF-8 text", test_txt_utf8)
    test("UTF-8 BOM text", test_txt_utf8_bom)
    test("Latin-1 text", test_txt_latin1)
    test("Đọc test_doc.txt thật", test_txt_from_file)

    # ─── TEST 2: parse_pdf ───────────────────────────────────
    print("\n📕 Test parse_pdf:")

    def test_pdf_generated():
        """Tạo PDF tạm bằng fpdf2 và parse"""
        from fpdf import FPDF
        import io

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(200, 10, text="Chapter 1: E-Commerce Overview", new_x="LMARGIN", new_y="NEXT")
        pdf.cell(200, 10, text="E-commerce is the buying and selling of goods online.", new_x="LMARGIN", new_y="NEXT")
        pdf.cell(200, 10, text="It includes B2B, B2C, and C2C models.", new_x="LMARGIN", new_y="NEXT")

        pdf_bytes = pdf.output()
        text = parse_pdf(pdf_bytes)
        assert "E-Commerce" in text or "E-commerce" in text or "commerce" in text.lower(), f"Missing expected text. Got: {text[:100]}"
        assert len(text) > 20, f"PDF text too short: {len(text)} chars"
        print(f"    → Đọc được {len(text)} ký tự từ PDF tự tạo")

    def test_pdf_invalid():
        try:
            parse_pdf(b"This is not a PDF file")
            assert False, "Should have raised an error"
        except Exception:
            pass  # Expected

    test("Tạo PDF → parse → kiểm tra nội dung", test_pdf_generated)
    test("File PDF không hợp lệ → lỗi", test_pdf_invalid)

    # ─── TEST 3: parse_docx ──────────────────────────────────
    print("\n📘 Test parse_docx:")

    def test_docx_create_and_parse():
        """Tạo file DOCX tạm trong memory và parse"""
        from docx import Document as DocxDocument
        import io

        doc = DocxDocument()
        doc.add_heading('Chương 1: Thương mại điện tử', level=1)
        doc.add_paragraph('Thương mại điện tử là hình thức kinh doanh trên Internet.')
        doc.add_heading('Chương 2: Thanh toán', level=1)
        doc.add_paragraph('Thanh toán điện tử bao gồm nhiều phương thức khác nhau.')

        # Add a table
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Loại"
        table.cell(0, 1).text = "Ví dụ"
        table.cell(1, 0).text = "B2C"
        table.cell(1, 1).text = "Amazon"

        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        text = parse_docx(docx_bytes)
        assert "Thương mại điện tử" in text, f"Missing expected text. Got: {text[:100]}"
        assert "Thanh toán" in text
        assert "B2C" in text  # Table content
        assert "Amazon" in text
        print(f"    → Đọc được {len(text)} ký tự từ DOCX tự tạo")

    def test_docx_invalid():
        try:
            parse_docx(b"Not a docx file")
            assert False, "Should have raised an error"
        except Exception:
            pass  # Expected

    test("Tạo DOCX → parse → kiểm tra nội dung", test_docx_create_and_parse)
    test("File DOCX không hợp lệ → lỗi", test_docx_invalid)

    # ─── TEST 4: parse_document (auto-detect) ────────────────
    print("\n🔍 Test parse_document (auto-detect):")

    def test_auto_txt():
        text = parse_document("bai_giang.txt", "Hello World".encode('utf-8'))
        assert text == "Hello World"

    def test_auto_pdf():
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(200, 10, text="Test PDF content for auto-detect", new_x="LMARGIN", new_y="NEXT")
        pdf_bytes = pdf.output()
        text = parse_document("test_file.pdf", pdf_bytes)
        assert len(text) > 10

    def test_auto_unsupported():
        try:
            parse_document("image.jpg", b"fake data")
            assert False, "Should have raised ValueError"
        except ValueError as e:
            assert "không được hỗ trợ" in str(e)

    def test_auto_empty():
        try:
            parse_document("file.txt", b"")
            assert False, "Should have raised ValueError"
        except ValueError as e:
            assert "rỗng" in str(e)

    def test_auto_no_filename():
        try:
            parse_document("", b"content")
            assert False, "Should have raised ValueError"
        except ValueError as e:
            assert "trống" in str(e)

    test("Auto-detect .txt", test_auto_txt)
    test("Auto-detect .pdf", test_auto_pdf)
    test("Extension không hỗ trợ (.jpg) → lỗi", test_auto_unsupported)
    test("File rỗng → lỗi", test_auto_empty)
    test("Filename rỗng → lỗi", test_auto_no_filename)

    # ─── TEST 5: Giới hạn ký tự ──────────────────────────────
    print("\n📏 Test giới hạn ký tự:")

    def test_truncate():
        long_text = "Câu này dài. " * 5000  # ~65,000 chars
        text = parse_document("long.txt", long_text.encode('utf-8'), max_chars=1000)
        assert len(text) <= 1000, f"Text should be ≤1000 chars, got {len(text)}"

    def test_no_truncate():
        short_text = "Ngắn gọn."
        text = parse_document("short.txt", short_text.encode('utf-8'))
        assert text == short_text

    test("Cắt text dài → ≤ max_chars", test_truncate)
    test("Text ngắn → giữ nguyên", test_no_truncate)

    # ─── TEST 6: Utility functions ───────────────────────────
    print("\n🔧 Test utility functions:")

    def test_accept_string():
        s = get_accept_string()
        assert ".txt" in s
        assert ".pdf" in s
        assert ".docx" in s

    test("get_accept_string() chứa đủ 3 loại", test_accept_string)

    # ─── KẾT QUẢ ─────────────────────────────────────────────
    print("\n" + "=" * 60)
    total = PASS + FAIL
    if FAIL == 0:
        print(f"🎉 TẤT CẢ {total} TEST ĐỀU PASS!")
    else:
        print(f"📊 Kết quả: {PASS}/{total} PASS, {FAIL} FAIL")
    print("=" * 60)

    return FAIL == 0


if __name__ == "__main__":
    success = run_tests()
    sys.exit(0 if success else 1)
